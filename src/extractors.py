"""Text extraction from office documents: .pdf, .docx, .eml, .xlsx, .csv, .txt, .md.

Used for the knowledge base (src/ingest_data.py) and for attachments (src/attachments.py), so a
document reads the same way wherever it comes from. Every extractor takes bytes, which lets e-mail
attachments go through the same code. A file that cannot be read raises ``ExtractionError``.
PDF pages without a text layer (scans) are read by OCR (src/ocr.py); their segments carry ``ocr``
and ``ocr_confidence``.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path

SUPPORTED = (".pdf", ".docx", ".eml", ".xlsx", ".csv", ".txt", ".md")
CSV_ROWS = 2000  # rows of a .csv put into text (the index, the prompt); the whole file stays readable for code


class ExtractionError(RuntimeError):
    pass


@dataclass
class Segment:
    """A piece of a document: the whole text, or one PDF page / one spreadsheet sheet."""

    text: str
    meta: dict = field(default_factory=dict)


def is_supported(name: str) -> bool:
    return Path(name).suffix.lower() in SUPPORTED


def extract(path: Path) -> list[Segment]:
    return extract_bytes(path.name, path.read_bytes())


def extract_bytes(name: str, data: bytes) -> list[Segment]:
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".pdf":
            segments = _pdf(data, name)
        elif suffix == ".docx":
            segments = [Segment(_docx(data))]
        elif suffix == ".eml":
            segments = [Segment(_eml(data))]
        elif suffix == ".xlsx":
            segments = _xlsx(data)
        elif suffix == ".csv":
            segments = _csv(data)
        elif suffix in (".txt", ".md"):
            segments = [Segment(_text(data))]
        else:
            raise ExtractionError(f"unsupported file type: {suffix or '(none)'}")
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"{name}: {type(exc).__name__}: {exc}") from exc
    segments = [s for s in segments if s.text.strip()]
    if not segments:
        raise ExtractionError(f"{name}: no text found")
    return segments


def as_text(segments: list[Segment]) -> str:
    parts = []
    for s in segments:
        label = f"[strona {s.meta['page']}]\n" if "page" in s.meta else (f"[arkusz {s.meta['sheet']}]\n" if "sheet" in s.meta else "")
        parts.append(label + s.text)
    return "\n\n".join(parts)


# --------------------------------------------------------------------------- formats

def _text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1250"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _pdf(data: bytes, name: str = "") -> list[Segment]:
    from pypdf import PdfReader

    from src import ocr
    from src.config import get_settings

    reader = PdfReader(io.BytesIO(data))
    segments = [Segment(page.extract_text() or "", {"page": i + 1}) for i, page in enumerate(reader.pages)]
    scanned = [s.meta["page"] for s in segments if ocr.needs_ocr(s.text)]
    if not scanned:
        return segments
    try:
        texts = ocr.read_pages(data, scanned)
    except ocr.OCRUnavailable as exc:
        if len(scanned) == len(segments):
            raise ExtractionError(f"{name}: no text found — scanned pages, OCR not available: {exc}") from exc
        if get_settings().ocr_enabled:
            print(f"⚠️ {name}: {len(scanned)} zeskanowanych stron bez tekstu — OCR nie zadziałał ({exc}); "
                  "po naprawie: `hapilotctl.py ingest --force`")
        return segments
    for s in segments:
        page = texts.get(s.meta["page"])
        if page is not None and not ocr.needs_ocr(page.text):
            s.text = page.text
            s.meta.update(ocr=True, ocr_confidence=page.confidence)
    return segments


def _docx(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, doc).text.strip()
            if text:
                lines.append(text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                cells = [c.text.strip() for c in row.cells]
                # merged cells repeat their text; keep each value once per row
                deduped = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
                if any(deduped):
                    lines.append(" | ".join(deduped))
    return "\n".join(lines)


class _HTMLText(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "table", "ul", "ol"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")
        elif tag == "td":
            self.parts.append(" | ")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip = max(0, self.skip - 1)
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    parser = _HTMLText()
    parser.feed(html)
    text = "".join(parser.parts)
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


def _eml(data: bytes) -> str:
    msg = BytesParser(policy=policy.default).parsebytes(data)
    lines = [f"{h}: {msg[h]}" for h in ("From", "To", "Cc", "Date", "Subject") if msg[h]]
    body = msg.get_body(preferencelist=("plain", "html"))
    if body is not None:
        content = body.get_content()
        lines += ["", html_to_text(content) if body.get_content_type() == "text/html" else content.strip()]
    attachments = [part for part in msg.iter_attachments() if part.get_filename()]
    if attachments:
        lines += ["", "Załączniki: " + ", ".join(p.get_filename() for p in attachments)]
        for part in attachments:
            name = part.get_filename()
            if is_supported(name) and not name.lower().endswith(".eml"):
                try:
                    text = as_text(extract_bytes(name, part.get_payload(decode=True) or b""))
                    lines += ["", f"--- Załącznik: {name} ---", text]
                except ExtractionError as exc:
                    lines += ["", f"--- Załącznik: {name} (nie udało się odczytać: {exc}) ---"]
    return "\n".join(lines).strip()


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _csv(data: bytes) -> list[Segment]:
    """Rows as " | "-separated lines (like a sheet), delimiter detected; at most ``CSV_ROWS`` rows, with a note."""
    import csv

    text = _text(data)
    try:
        delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    rows, total = [], 0
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        cells = [c.strip() for c in row]
        while cells and not cells[-1]:
            cells.pop()
        if not any(cells):
            continue
        total += 1
        if len(rows) < CSV_ROWS:
            rows.append(" | ".join(cells))
    if not rows:
        return []
    body = "\n".join(rows)
    if total > len(rows):
        body += f"\n[… pokazano {len(rows)} z {total} wierszy — cały plik może przeanalizować kod Python]"
    return [Segment(body, {"header": rows[0]})]


def _xlsx(data: bytes) -> list[Segment]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    segments = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(v) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            segments.append(Segment("\n".join(rows), {"sheet": ws.title, "header": rows[0]}))
    wb.close()
    return segments
