"""The structure of data files (.xlsx, .csv) for the code generator: sheets, row counts, columns and first rows.

The coder used to see only file names and guessed the layout. Seen on the pilot: an Excel file whose every row was
one cell holding a whole CSV line (``4,"Poziom","0.012","m","2026-07-13 12:30:00"``) — five attempts, five empty
charts. Such sheets are recognised here and the preview says how to read them.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from pathlib import Path

DATA_SUFFIXES = (".xlsx", ".csv")
SHOWN_ROWS = 5
MAX_SHEETS = 3
_LINE = 240
_DELIMITERS = ",;\t|"


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _line(cells: list[str]) -> str:
    text = " | ".join(cells)
    return text if len(text) <= _LINE else text[:_LINE] + " …"


def _is_number(text: str) -> bool:
    try:
        float(text.replace(",", "."))
        return True
    except ValueError:
        return False


def csv_in_cells(lines: list[str]) -> tuple[str, int] | None:
    """(delimiter, fields) when single-cell rows are whole CSV lines with the same number of fields, else None."""
    lines = [ln for ln in lines if ln.strip()][:20]
    if len(lines) < 2:
        return None
    for delimiter in _DELIMITERS:
        counts = [len(next(csv.reader([ln], delimiter=delimiter))) for ln in lines]
        if min(counts) >= 2 and max(counts) == min(counts):
            return delimiter, counts[0]
    return None


def _xlsx(path: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        out = []
        for ws in wb.worksheets[:MAX_SHEETS]:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [_cell(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append(cells)
                if len(rows) >= 20:
                    break
            size = f"ok. {ws.max_row} wierszy" if ws.max_row else "liczba wierszy nieznana"
            title = f"Plik {path.name} — arkusz „{ws.title}”: {size}"
            if not rows:
                out.append(f"{title}, pusty.")
                continue
            single = [r[0] for r in rows if len(r) == 1]
            found = csv_in_cells(single) if len(single) >= 0.8 * len(rows) else None
            if found:
                delimiter, fields = found
                first = next(csv.reader([single[0]], delimiter=delimiter))
                header = not any(_is_number(f) for f in first)
                shown = repr(delimiter)[1:-1]
                out.append(
                    f"{title}. UWAGA: każdy wiersz to JEDNA komórka z całą linią CSV (separator \"{shown}\", {fields} pól, "
                    f"{'pierwszy wiersz to nagłówek' if header else 'bez nagłówka — pierwszy wiersz to już dane'}). "
                    "Przykładowe wiersze:\n" + "\n".join("  " + _line([s]) for s in single[:SHOWN_ROWS])
                    + f"\n  Wczytaj tak: raw = pd.read_excel(\"{path.name}\", header=None, dtype=str)[0].dropna()\n"
                    f"               df = pd.read_csv(io.StringIO(\"\\n\".join(raw)), sep=\"{shown}\", "
                    f"header={'0' if header else 'None'})  # import io")
                continue
            out.append(f"{title}; kolumny (1. wiersz): {_line(rows[0])}\n  dalsze wiersze:\n"
                       + "\n".join("  " + _line(r) for r in rows[1:1 + SHOWN_ROWS]))
        if len(wb.worksheets) > MAX_SHEETS:
            out.append(f"Plik {path.name}: jeszcze {len(wb.worksheets) - MAX_SHEETS} arkusz(e) — "
                       f"nazwy: {', '.join(ws.title for ws in wb.worksheets[MAX_SHEETS:])}")
        return out
    finally:
        wb.close()


def _csv(path: Path) -> list[str]:
    raw = path.read_bytes()
    sample_bytes = raw[:65536]
    for encoding in ("utf-8-sig", "cp1250"):
        try:
            sample = sample_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        sample, encoding = sample_bytes.decode("latin-1"), "latin-1"
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=_DELIMITERS).delimiter
    except csv.Error:
        delimiter = ","
    rows = [r for r in csv.reader(io.StringIO(sample), delimiter=delimiter) if any(c.strip() for c in r)][:1 + SHOWN_ROWS]
    lines = raw.count(b"\n") + (0 if raw.endswith(b"\n") else 1)
    if not rows:
        return [f"Plik {path.name}: pusty."]
    header = not any(_is_number(c) for c in rows[0] if c.strip())
    shown = repr(delimiter)[1:-1]
    return [f"Plik {path.name}: ok. {lines} wierszy, separator \"{shown}\", kodowanie {encoding}, "
            f"{'pierwszy wiersz to nagłówek' if header else 'bez nagłówka'}; pierwsze wiersze:\n"
            + "\n".join("  " + _line([c.strip() for c in r]) for r in rows)
            + f"\n  Wczytaj tak: df = pd.read_csv(\"{path.name}\", sep=\"{shown}\", encoding=\"{encoding}\", "
            f"header={'0' if header else 'None'})"]


def preview(path: Path) -> str:
    """The structure of one data file, in Polish, for the coder's prompt ("" for other files)."""
    try:
        if path.suffix.lower() == ".xlsx":
            return "\n".join(_xlsx(path))
        if path.suffix.lower() == ".csv":
            return "\n".join(_csv(path))
    except Exception as exc:  # a damaged file: say so, the code will meet the same error
        return f"Plik {path.name}: nie udało się odczytać struktury ({type(exc).__name__}: {exc})"
    return ""
