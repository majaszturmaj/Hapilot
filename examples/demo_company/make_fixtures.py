#!/usr/bin/env python3
"""(Re)generate the binary demo documents of the fictional company Przykładowa Sp. z o.o.

    .venv/bin/python examples/demo_company/make_fixtures.py [--uploads-only]

Writes docs/umowa_serwisowa_7_2026.docx, docs/cennik_uslug_2026.pdf (via LibreOffice, from a
temporary .docx), docs/zgloszenie_awarii_2026-10-12.eml, docs/rejestr_reklamacji_2026.xlsx, the
evaluation attachment tests/eval/attachments/rozliczenie_delegacji.xlsx, and the files to upload in the
demo: uploads/projekt_umowy_sprzatanie.docx and uploads/raport_kwartalny_q3_2026.pdf (via LibreOffice).
All content is invented; e-mail addresses use example.com.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from datetime import date
from email.message import EmailMessage
from pathlib import Path

from docx import Document
from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
DOCS = HERE / "docs"
UPLOADS = HERE / "uploads"
ROOT = HERE.parent.parent
COMPANY = "Przykładowa Sp. z o.o., ul. Zmyślona 7, 50-001 Wrocław"


def service_agreement() -> None:
    doc = Document()
    doc.add_heading("Umowa serwisowa nr 7/2026", level=1)
    doc.add_paragraph(f"Zawarta 2 stycznia 2026 r. między {COMPANY} (Zamawiający) a Serwis-Wymyślony Sp. z o.o. "
                      "(Wykonawca).")
    doc.add_paragraph("§ 1. Przedmiot umowy: serwis serwerów i drukarek Zamawiającego w siedzibie we Wrocławiu.")
    doc.add_paragraph("§ 2. Zakres i ceny usług:")
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["Usługa", "Cena netto", "Częstotliwość"]):
        table.rows[0].cells[i].text = h
    for row in [("Przegląd serwerowni", "1 200 zł", "raz na kwartał"), ("Przegląd drukarek", "450 zł", "raz na pół roku"),
                ("Interwencja awaryjna", "300 zł za godzinę", "na zgłoszenie")]:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    doc.add_paragraph("§ 3. Czas reakcji: Wykonawca podejmuje interwencję awaryjną w ciągu 8 godzin roboczych od zgłoszenia "
                      "na adres serwis@example.com.")
    doc.add_paragraph("§ 4. Kara umowna: za każdą rozpoczętą godzinę opóźnienia ponad czas reakcji Wykonawca zapłaci 100 zł.")
    doc.add_paragraph("§ 5. Umowa obowiązuje do 31 grudnia 2026 r. z miesięcznym okresem wypowiedzenia.")
    doc.save(DOCS / "umowa_serwisowa_7_2026.docx")


def price_list_pdf() -> None:
    soffice = shutil.which("soffice")
    if not soffice:
        print("LibreOffice (soffice) not found — skipping the PDF")
        return
    doc = Document()
    doc.add_heading("Cennik usług biurowych 2026", level=1)
    doc.add_paragraph(f"{COMPANY}. Ceny netto, obowiązują od 1 lutego 2026 r. Do cen doliczamy 23% VAT.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Usługa", "Cena netto"
    for name, price in [("Tłumaczenie strony A4 (polski–angielski)", "60 zł"), ("Skład raportu (do 20 stron)", "350 zł"),
                        ("Archiwizacja dokumentów (karton)", "25 zł"), ("Digitalizacja (100 stron)", "40 zł")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = name, price
    doc.add_paragraph("Zamówienia powyżej 2 000 zł netto realizujemy z rabatem 5%.")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "cennik_uslug_2026.docx"
        doc.save(src)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)], check=True,
                       capture_output=True, timeout=120)
        shutil.copy(Path(tmp) / "cennik_uslug_2026.pdf", DOCS / "cennik_uslug_2026.pdf")


def failure_report_eml() -> None:
    msg = EmailMessage()
    msg["From"] = "Anna Nowak <anna.nowak@example.com>"
    msg["To"] = "helpdesk@example.com"
    msg["Date"] = "Mon, 12 Oct 2026 08:47:00 +0200"
    msg["Subject"] = "Awaria drukarki w Dziale Sprzedaży"
    msg.set_content("Dzień dobry,\n\nod rana drukarka sieciowa DR-2 w Dziale Sprzedaży (pokój 14) nie drukuje — "
                    "wyświetla błąd E-52. Restart nie pomógł. Prosimy o interwencję serwisu.\n\n"
                    "Pozdrawiam\nAnna Nowak\nDział Sprzedaży")
    msg.add_attachment("2026-10-12 08:15 DR-2 error E-52 (fuser)\n2026-10-12 08:31 DR-2 error E-52 (fuser)\n",
                       filename="dziennik_drukarki.txt")
    (DOCS / "zgloszenie_awarii_2026-10-12.eml").write_bytes(msg.as_bytes())


def complaints_register_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rejestr 2026"
    ws.append(["Nr", "Data wpływu", "Klient", "Kwota (zł)", "Decyzja", "Data odpowiedzi"])
    rows = [("R-01", date(2026, 3, 2), "Hurtownia Wymyślona", 1500, "uznana", date(2026, 3, 12)),
            ("R-02", date(2026, 3, 9), "Biuro Testowe", 320, "odrzucona", date(2026, 3, 20)),
            ("R-03", date(2026, 3, 23), "Hurtownia Wymyślona", 780, "uznana", date(2026, 4, 3)),
            ("R-04", date(2026, 4, 14), "Drukarnia Przykład", 2100, "uznana", date(2026, 4, 30)),
            ("R-05", date(2026, 5, 6), "Biuro Testowe", 95, "odrzucona", date(2026, 5, 15))]
    for r in rows:
        ws.append(list(r))
    wb.save(DOCS / "rejestr_reklamacji_2026.xlsx")


def delegation_attachment_xlsx() -> None:
    out = ROOT / "tests" / "eval" / "attachments"
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Delegacje październik"
    ws.append(["Pracownik", "Cel", "Koszt (zł)"])
    for r in [("Jan Kowalski", "Poznań", 412.50), ("Anna Nowak", "Kraków", 389.00), ("Piotr Testowy", "Gdańsk", 655.25),
              ("Ewa Przykładowa", "Łódź", 243.00)]:
        ws.append(list(r))
    wb.save(out / "rozliczenie_delegacji.xlsx")


def _to_pdf(doc, out: Path) -> None:
    soffice = shutil.which("soffice")
    if not soffice:
        print(f"LibreOffice (soffice) not found — skipping {out.name}")
        return
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / (out.stem + ".docx")
        doc.save(src)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)], check=True,
                       capture_output=True, timeout=120)
        shutil.copy(Path(tmp) / (out.stem + ".pdf"), out)


def cleaning_contract_draft() -> None:
    """A draft with deliberate typos and a 14-day payment term — the demo asks Hapilot to fix both."""
    doc = Document()
    doc.add_heading("Projekt umowy o świadczenie usług sprzątania", level=1)
    doc.add_paragraph(f"Zawarta w dniu … 2026 r. we Wrocławiu pomiędzy {COMPANY}, zwaną dalej Zamawiającym, "
                      "a Czysto-Wymyślone Sp. z o.o., ul. Fikcyjna 3, 50-003 Wrocław, zwaną dalej Wykonwacą.")
    doc.add_heading("§ 1. Przedmiot umowy", level=2)
    doc.add_paragraph("Wykonwaca zobowiązuje się do sprzątania pomieszczeń biurowych Zamawiającego o łącznej powierzchni "
                      "420 m², ktróre znajdują się na parterze i pierwszym piętrze budynku.")
    doc.add_heading("§ 2. Zakres prac", level=2)
    for item in ["odkurzanie i mycie podłóg — codziennie od poniedziałku do piątku",
                 "opróżnianie koszy i segregacja odpadów — codziennie",
                 "mycie okien — raz na kwartał",
                 "sprzątanie kuchni i łazienek wraz z uzupełnianiem środków higienicznych — codzienie"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("§ 3. Wynagrodzenie", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Usługa", "Cena netto"
    for name, price in [("Sprzątanie codzienne (miesięcznie)", "3 800 zł"), ("Mycie okien (jednorazowo)", "900 zł")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = name, price
    doc.add_paragraph("Zamawiający zapłaci wynagrodzenie przelewem w terminie 14 dni od dnia otrzymania prawidłowo "
                      "wystawionej faktury.")
    doc.add_heading("§ 4. Postanowienia końcowe", level=2)
    doc.add_paragraph("Umowa zostaje zawarta na czas nieokreślony z jednomiesięcznym okresem wypowiedzenia. "
                      "Zmiany umowy wymagają formy pisemnej pod rygorem nieważności.")
    doc.save(UPLOADS / "projekt_umowy_sprzatanie.docx")


MONTHS = [("lipiec", 1_520_000, 1_210_000, 41, 3), ("sierpień", 1_440_000, 1_185_000, 40, 5),
          ("wrzesień", 1_900_000, 1_330_000, 43, 2)]
DEPARTMENTS = [("Dział Sprzedaży", 14, "pozyskał 9 nowych klientów biznesowych, głównie z branży logistycznej"),
               ("Dział Obsługi Klienta", 8, "obsłużył 1 284 zgłoszenia, z czego 93% w ciągu jednego dnia roboczego"),
               ("Dział IT", 5, "przeprowadził migrację poczty i wymienił 12 stanowisk komputerowych"),
               ("Dział Kadr", 3, "przeprowadził 6 rekrutacji i wdrożył nowy system ewidencji czasu pracy"),
               ("Dział Jakości", 4, "rozpatrzył 11 reklamacji, z których 7 uznano"),
               ("Dział Finansów", 6, "zamknął półrocze bez uwag audytora i rozpoczął przegląd umów z dostawcami")]


def quarterly_report_pdf() -> None:
    """A report long enough (~25 000 characters) to be indexed instead of inlined and summarised part by part."""
    doc = Document()
    doc.add_heading("Raport kwartalny — III kwartał 2026", level=1)
    doc.add_paragraph(f"{COMPANY}. Opracowanie: Dział Finansów, październik 2026 r. Dokument wewnętrzny.")
    doc.add_heading("1. Podsumowanie", level=2)
    doc.add_paragraph("W III kwartale 2026 r. przychody ze sprzedaży wyniosły 4 860 000 zł, o 11% więcej niż w II kwartale. "
                      "Koszty operacyjne wyniosły 3 725 000 zł, a zysk operacyjny 1 135 000 zł. Zatrudnienie na koniec "
                      "kwartału: 43 osoby. Najważniejsze wydarzenia: otwarcie biura obsługi klienta w Opolu (1 września), "
                      "migracja poczty do nowego systemu oraz podpisanie umowy ramowej z Hurtownią Wymyśloną.")
    doc.add_heading("2. Sprzedaż i koszty w poszczególnych miesiącach", level=2)
    table = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["Miesiąc", "Przychody (zł)", "Koszty (zł)", "Zatrudnienie"]):
        table.rows[0].cells[i].text = h
    for month, revenue, cost, staff, _ in MONTHS:
        cells = table.add_row().cells
        for i, v in enumerate([month, f"{revenue:,}".replace(",", " "), f"{cost:,}".replace(",", " "), str(staff)]):
            cells[i].text = v
    for month, revenue, cost, staff, incidents in MONTHS:
        margin = round(100 * (revenue - cost) / revenue, 1)
        for k in range(6):
            doc.add_paragraph(
                f"{month.capitalize()} — część {k + 1}. Przychody w miesiącu wyniosły {revenue:,} zł".replace(",", " ")
                + f", koszty {cost:,} zł".replace(",", " ")
                + f", co daje marżę operacyjną {margin}%. Zatrudnienie na koniec miesiąca: {staff} osób. "
                "Na wynik wpłynęły przede wszystkim zamówienia stałych klientów oraz sezonowe usługi archiwizacji. "
                f"W tym miesiącu odnotowano {incidents} incydenty informatyczne, wszystkie usunięte w czasie "
                "określonym w polityce bezpieczeństwa. Dział Finansów zwraca uwagę na rosnący udział kosztów "
                "energii i najmu, które w kolejnych miesiącach będą monitorowane w comiesięcznych zestawieniach. "
                "Należności przeterminowane pozostają na poziomie poniżej 4% przychodów.")
    doc.add_heading("3. Działy", level=2)
    for name, people, result in DEPARTMENTS:
        doc.add_heading(name, level=3)
        for k in range(4):
            doc.add_paragraph(
                f"{name} (zatrudnienie: {people} osób) w III kwartale {result}. "
                + ("Kierownik działu ocenia realizację planu kwartalnego jako zgodną z założeniami. " if k == 0 else
                   "Do najważniejszych wyzwań należały nieobecności urlopowe w sierpniu i wdrażanie nowych pracowników. "
                   if k == 1 else
                   "Dział zgłosił potrzebę doposażenia stanowisk pracy i doszkolenia zespołu z nowych narzędzi. "
                   if k == 2 else
                   "Na IV kwartał dział planuje kontynuację bieżących zadań oraz udział w przeglądzie procesów. ")
                + "Szczegółowe dane liczbowe znajdują się w zestawieniach działowych przechowywanych w systemie "
                "dokumentów. Wnioski z kwartału zostaną omówione na spotkaniu kierowników w listopadzie.")
    doc.add_heading("4. Ryzyka", level=2)
    for risk in ["wzrost cen energii o ponad 15% od stycznia 2027 r.", "odejście dwóch kluczowych klientów z branży "
                 "wydawniczej", "opóźnienia w dostawach sprzętu komputerowego", "rotacja pracowników w Dziale Obsługi Klienta"]:
        doc.add_paragraph(risk, style="List Bullet")
        doc.add_paragraph("Ryzyko jest monitorowane przez zarząd raz w miesiącu; działania ograniczające ryzyko opisano "
                          "w planie ciągłości działania. Właściciel ryzyka raportuje stan na spotkaniach kierowników.")
    doc.add_heading("5. Plan na IV kwartał", level=2)
    doc.add_paragraph("Cele na IV kwartał 2026 r.: przychody co najmniej 5 100 000 zł, uruchomienie samoobsługowego "
                      "portalu klienta do 30 listopada, zakończenie wymiany drukarek w Dziale Sprzedaży oraz przeprowadzenie "
                      "szkolenia z bezpieczeństwa informacji dla wszystkich pracowników do 15 grudnia.")
    _to_pdf(doc, UPLOADS / "raport_kwartalny_q3_2026.pdf")


if __name__ == "__main__":
    import sys

    if "--uploads-only" in sys.argv:  # leave docs/ untouched (its files are part of the development tasks)
        UPLOADS.mkdir(parents=True, exist_ok=True)
        cleaning_contract_draft()
        quarterly_report_pdf()
        raise SystemExit(0)
    DOCS.mkdir(parents=True, exist_ok=True)
    service_agreement()
    price_list_pdf()
    failure_report_eml()
    complaints_register_xlsx()
    delegation_attachment_xlsx()
    UPLOADS.mkdir(parents=True, exist_ok=True)
    cleaning_contract_draft()
    quarterly_report_pdf()
    print("fixtures written to", DOCS, "and", UPLOADS)
