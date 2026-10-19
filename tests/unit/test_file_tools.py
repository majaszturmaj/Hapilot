"""Uploaded files: summary (short and long), editing (.docx, .txt; refusals), adding to the knowledge base."""

import json

import pytest
from langchain_core.messages import HumanMessage

from src import attachments, run_context


def plan(*steps):
    return json.dumps({"task_name": "t", "complexity": "low",
                       "steps": [{"id": i + 1, "action": a, "tool": t, "input": a} for i, (t, a) in enumerate(steps)]})


@pytest.fixture(autouse=True)
def _no_vector_index(monkeypatch):
    monkeypatch.setattr(attachments, "_index", lambda texts, path, name: None)


def attach(tmp_path, sid, text, *files):
    """A user message with files attached the way the UI does it."""
    with run_context.use_context(session_id=sid):
        prepared = attachments.prepare(list(files), index_path=tmp_path / "idx")
    return HumanMessage(content=text + prepared.message_block)


def _docx(path):
    from docx import Document

    doc = Document()
    doc.add_heading("Umowa serwisowa", level=1)
    doc.add_paragraph("Termin płatności wynosi 14 dni.")
    doc.add_paragraph("Przegląd raz w miesiącu", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Usługa", "Cena"
    table.cell(1, 0).text, table.cell(1, 1).text = "Przegląd", "300 zł"
    doc.save(path)
    return path


def last_text(state):
    return state["messages"][-1].content


# --------------------------------------------------------------------------- finding the files

def test_pick_prefers_named_files_then_the_latest_upload(tmp_path):
    (tmp_path / "raport.txt").write_text("a")
    (tmp_path / "umowa.txt").write_text("b")
    messages = [attach(tmp_path, "pick", "pierwszy", tmp_path / "raport.txt"),
                attach(tmp_path, "pick", "drugi", tmp_path / "umowa.txt")]
    assert attachments.session_attachments(messages) == (["umowa.txt"], ["raport.txt", "umowa.txt"])
    assert attachments.pick("Streść raport", messages) == ["raport.txt"]
    assert attachments.pick("Streść to", messages) == ["umowa.txt"]
    assert attachments.pick("cokolwiek", [HumanMessage(content="bez plików")]) == []


# --------------------------------------------------------------------------- summary

def test_short_file_is_summarised_from_its_whole_text(graph, invoke, llm, tmp_path):
    f = tmp_path / "notatka.txt"
    f.write_text("Spotkanie zarządu 3 listopada. Budżet IT: 120 000 zł. Odpowiada Anna Nowak.", encoding="utf-8")
    llm.on("task_planner", plan(("file_summary", "Streść załączoną notatkę")))
    llm.on("file_assistant", "Zarząd 3 listopada; budżet IT 120 000 zł (A. Nowak).")
    state = invoke(graph, [attach(tmp_path, "sum-short", "Streść ten plik", f)], "sum-short")
    assert last_text(state) == "Zarząd 3 listopada; budżet IT 120 000 zł (A. Nowak)."
    prompt = llm.prompts_for("file_assistant")[0]
    assert "Budżet IT: 120 000 zł" in prompt and "Polecenie użytkownika: Streść ten plik" in prompt


def test_long_file_is_summarised_part_by_part(graph, invoke, llm, tmp_path):
    f = tmp_path / "raport.txt"
    f.write_text("\n\n".join(f"Rozdział {i}. " + "Treść rozdziału o wynikach sprzedaży. " * 40 for i in range(20)),
                 encoding="utf-8")
    llm.on("task_planner", plan(("file_summary", "Streść raport")))
    llm.on("file_assistant", lambda msgs: "całość" if "Notatki z kolejnych" in msgs[-1].content else "notatka")
    state = invoke(graph, [attach(tmp_path, "sum-long", "Streść załączony raport", f)], "sum-long")
    prompts = llm.prompts_for("file_assistant")
    assert last_text(state) == "całość"
    assert len(prompts) >= 3 and "Notatki z kolejnych fragmentów pliku raport.txt" in prompts[-1]
    assert prompts[-1].count("[Fragment ") == len(prompts) - 1


def test_very_long_file_summary_says_what_it_covered(llm):
    from extensions.file_tools.nodes import summarize_text

    llm.on("file_assistant", "x")
    from src.llm_config import get_llm

    model, system = get_llm("file_assistant")
    text = "Akapit. " * 5000
    result = summarize_text(text, "duzy.txt", "streść", model, system, {"single_pass_chars": 1000, "part_chars": 1000, "max_parts": 2})
    assert result.startswith("Uwaga: plik jest bardzo długi — streszczenie obejmuje pierwsze 2 z ")


def test_no_attachment_is_explained(graph, invoke, llm):
    llm.on("task_planner", plan(("file_summary", "Streść plik")))
    state = invoke(graph, "Streść ten plik", "sum-none")
    assert "Nie widzę w tej rozmowie załączonego pliku" in last_text(state)
    assert not llm.prompts_for("file_assistant")


# --------------------------------------------------------------------------- edit

def test_docx_round_trip_keeps_structure(tmp_path):
    from docx import Document

    from extensions.file_tools.nodes import docx_to_markup, markup_to_docx

    src = _docx(tmp_path / "umowa.docx")
    markup = docx_to_markup(src)
    assert markup == ("# Umowa serwisowa\nTermin płatności wynosi 14 dni.\n- Przegląd raz w miesiącu\n\n"
                      "| Usługa | Cena |\n| Przegląd | 300 zł |\n")
    out = tmp_path / "out.docx"
    markup_to_docx(markup.replace("14 dni", "30 dni") + "\n| --- | --- |\nNowy akapit.\n", out, template=src)
    doc = Document(str(out))
    paras = [(p.style.name, p.text) for p in doc.paragraphs]
    assert paras == [("Heading 1", "Umowa serwisowa"), ("Normal", "Termin płatności wynosi 30 dni."),
                     ("List Bullet", "Przegląd raz w miesiącu"), ("Normal", "Nowy akapit.")]
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [["Usługa", "Cena"], ["Przegląd", "300 zł"]]


def test_edit_of_attached_docx_produces_a_new_word_file(graph, invoke, llm, tmp_path):
    from docx import Document

    src = _docx(tmp_path / "umowa.docx")
    llm.on("task_planner", plan(("file_edit", "Zmień termin płatności na 30 dni")))
    llm.on("file_assistant", lambda msgs: "```\n" + msgs[-1].content.split("<<<\n")[1].split("\n>>>")[0]
           .replace("14 dni", "30 dni") + "\n```")
    state = invoke(graph, [attach(tmp_path, "edit-docx", "Zmień termin płatności na 30 dni", src)], "edit-docx")
    assert "Wygenerowano plik: umowa_zmieniony.docx" in last_text(state)
    with run_context.use_context(session_id="edit-docx"):
        out = run_context.workspace_file("umowa_zmieniony.docx")
        original = run_context.workspace_file("umowa.docx")
    assert "Termin płatności wynosi 30 dni." in [p.text for p in Document(str(out)).paragraphs]
    assert "Termin płatności wynosi 14 dni." in [p.text for p in Document(str(original)).paragraphs]  # untouched


def test_edit_of_text_file_keeps_format_and_names_the_copy_uniquely(llm, tmp_path, session):
    from extensions.file_tools.nodes import edit_file
    from src.llm_config import get_llm

    f = run_context.workspace_dir() / "notatka.txt"
    f.write_text("Spotkanie w piątek.", encoding="utf-8")
    (run_context.workspace_dir() / "notatka_zmieniony.txt").write_text("starsza wersja")
    llm.on("file_assistant", "Spotkanie w czwartek.")
    model, system = get_llm("file_assistant")
    out = edit_file(f, "zmień dzień na czwartek", model, system, 20000)
    assert out.name == "notatka_zmieniony_1.txt" and out.read_text(encoding="utf-8") == "Spotkanie w czwartek.\n"


@pytest.mark.parametrize("name, content, message", [
    ("dane.xlsx", None, "edycja plików .xlsx nie jest obsługiwana"),
    ("dlugi.txt", "x" * 30000, "więcej niż limit edycji"),
])
def test_edit_refusals(graph, invoke, llm, tmp_path, name, content, message):
    f = tmp_path / name
    if content is None:
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.append(["a", 1])
        wb.save(f)
    else:
        f.write_text(content)
    llm.on("task_planner", plan(("file_edit", "Popraw plik")))
    state = invoke(graph, [attach(tmp_path, f"edit-{name}", "Popraw ten plik", f)], f"edit-{name}")
    assert message in "\n".join(m.content for m in state["messages"])
    assert not llm.prompts_for("file_assistant")


def test_edit_with_several_files_asks_which(graph, invoke, llm, tmp_path):
    (tmp_path / "a.txt").write_text("a")
    (tmp_path / "b.txt").write_text("b")
    llm.on("task_planner", plan(("file_edit", "Popraw plik")))
    state = invoke(graph, [attach(tmp_path, "edit-two", "Popraw literówki", tmp_path / "a.txt", tmp_path / "b.txt")], "edit-two")
    assert "Napisz, który mam zmienić" in last_text(state)


# --------------------------------------------------------------------------- knowledge base

@pytest.fixture
def added(monkeypatch):
    import src.ingest_data as ingest_data

    calls = []

    def fake_add(paths, **kw):
        calls.append(([p.name for p in paths], kw))
        return {"added": [f"dodane/{p.name}" for p in paths], "chunks": 3, "skipped": []}

    monkeypatch.setattr(ingest_data, "add_documents", fake_add)
    return calls


def test_adding_to_knowledge_waits_for_confirmation(graph, invoke, llm, tmp_path, added):
    f = tmp_path / "parking.txt"
    f.write_text("Parking dla gości: strefa B.", encoding="utf-8")
    llm.on("task_planner", plan(("file_to_knowledge", "Dodaj plik do bazy wiedzy")))
    first = invoke(graph, [attach(tmp_path, "kb-ask", "Dodaj ten plik do bazy wiedzy", f)], "kb-ask", approval_policy="ask")
    request = first["hitl_request"]
    assert request["type"] == "confirm" and "parking.txt" in request["content"] and added == []

    response = {**request["continuation"], "continuation_type": "resume_node", "confirmed": True}
    second = invoke(graph, first["messages"], "kb-ask", hitl_response=response)
    assert "Dodano do bazy wiedzy: dodane/parking.txt (3 fragm.)" in last_text(second)
    from src.config import get_settings

    names, kw = added[0]
    assert names == ["parking.txt"] and kw["added_dir"] == get_settings().added_docs_dir
    assert kw["collection"] == get_settings().collection


def test_declined_addition_changes_nothing(graph, invoke, llm, tmp_path, added):
    f = tmp_path / "parking.txt"
    f.write_text("Parking.", encoding="utf-8")
    llm.on("task_planner", plan(("file_to_knowledge", "Dodaj plik do bazy wiedzy")))
    first = invoke(graph, [attach(tmp_path, "kb-no", "Dodaj do bazy wiedzy", f)], "kb-no")
    response = {**first["hitl_request"]["continuation"], "continuation_type": "resume_node", "confirmed": False}
    second = invoke(graph, first["messages"], "kb-no", hitl_response=response)
    assert last_text(second) == "Rozumiem — nie dodałem plików do bazy wiedzy." and added == []


def test_auto_policy_adds_directly_but_evaluation_runs_never_do(graph, invoke, llm, tmp_path, added):
    f = tmp_path / "parking.txt"
    f.write_text("Parking.", encoding="utf-8")
    llm.on("task_planner", plan(("file_to_knowledge", "Dodaj plik do bazy wiedzy")))
    state = invoke(graph, [attach(tmp_path, "kb-auto", "Dodaj do bazy wiedzy", f)], "kb-auto", approval_policy="auto")
    assert "Dodano do bazy wiedzy" in last_text(state) and len(added) == 1

    from langchain_core.messages import HumanMessage as Human

    msg = attach(tmp_path, "kb-eval", "Dodaj do bazy wiedzy", f)
    config = run_context.make_config("kb-eval", approval_policy="auto", kb_path=str(tmp_path / "eval_idx"), kb_collection="eval_x")
    inputs = {"messages": [Human(content=msg.content)], "next_step": "", "execution_status": "", "retry_count": 0,
              "task_plan": {}, "task_id": "", "is_complex_task": False, "current_step": None, "generated_code": None,
              "last_generated_file": "", "hitl_request": None, "hitl_response": None, "halt": False}
    state = graph.invoke(inputs, config=config)
    assert "baza wiedzy jest stała" in last_text(state) and len(added) == 1


# --------------------------------------------------------------------------- planner guards with attachments

def test_edit_request_with_new_attachment_is_not_forced_to_modify_document(graph, invoke, llm, tmp_path, docker):
    src = _docx(tmp_path / "pismo.docx")
    llm.on("task_planner", plan(("file_edit", "Popraw załączone pismo")))
    llm.on("file_assistant", lambda msgs: msgs[-1].content.split("<<<\n")[1].split("\n>>>")[0])
    state = invoke(graph, [attach(tmp_path, "guard-att", "Popraw to pismo", src)], "guard-att",
                   last_file="letter_20260101_000000.tex")  # a PDF was generated earlier in the session
    assert [s["tool"] for s in state["task_plan"]["steps"]] == ["file_edit"]


def test_guards_read_the_users_words_not_the_attached_text(graph, invoke, llm, tmp_path):
    f = tmp_path / "zadania.txt"
    f.write_text("Na jutro: przygotuj raport w PDF dla zarządu.", encoding="utf-8")
    llm.on("task_planner", plan(("file_summary", "Streść plik")))
    llm.on("file_assistant", "Jedno zadanie: raport PDF dla zarządu.")
    state = invoke(graph, [attach(tmp_path, "guard-words", "Streść ten plik", f)], "guard-words")
    assert [s["tool"] for s in state["task_plan"]["steps"]] == ["file_summary"]  # no template_pdf step added


def test_modify_document_next_to_file_edit_is_dropped(graph, invoke, llm, tmp_path):
    # Bielik planned file_edit + modify_document for an attached contract (integration run 2026-10-19)
    f = tmp_path / "umowa.txt"
    f.write_text("Termin płatności 14 dni.", encoding="utf-8")
    llm.on("task_planner", plan(("file_edit", "Popraw umowę"), ("modify_document", "Zmień termin")))
    llm.on("file_assistant", "Termin płatności 30 dni.")
    state = invoke(graph, [attach(tmp_path, "guard-mod", "Zmień termin płatności na 30 dni", f)], "guard-mod")
    assert [s["tool"] for s in state["task_plan"]["steps"]] == ["file_edit"]
    assert "Wygenerowano plik: umowa_zmieniony.txt" in last_text(state)
