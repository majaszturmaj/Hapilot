"""Text extraction for every supported format (fixtures are built in the test; no binary files needed)."""

import io
from datetime import date
from email.message import EmailMessage

import pytest

from src.extractors import ExtractionError, as_text, extract_bytes, html_to_text


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Umowa serwisowa nr 7/2026")
    table = doc.add_table(rows=2, cols=3)
    for i, value in enumerate(["Usługa", "Cena netto", "Termin"]):
        table.rows[0].cells[i].text = value
    for i, value in enumerate(["Przegląd", "1200 zł", "co kwartał"]):
        table.rows[1].cells[i].text = value
    doc.add_paragraph("Zażółć gęślą jaźń — po tabeli.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Reklamacje"
    ws.append(["Nr", "Data wpływu", "Kwota", "Status"])
    ws.append(["R-1", date(2026, 3, 2), 1500.0, "uznana"])
    ws.append(["R-2", date(2026, 3, 9), 320.5, "odrzucona"])
    ws.append([None, None, None, None])
    ws2 = wb.create_sheet("Puste")
    ws2.append([None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_docx_keeps_order_of_paragraphs_and_tables():
    text = as_text(extract_bytes("umowa.docx", _docx_bytes()))
    lines = text.splitlines()
    assert lines[0] == "Umowa serwisowa nr 7/2026"
    assert lines[1] == "Usługa | Cena netto | Termin" and lines[2] == "Przegląd | 1200 zł | co kwartał"
    assert lines[3].startswith("Zażółć gęślą jaźń")


def test_xlsx_rows_per_sheet_with_values_not_formulas():
    segments = extract_bytes("rejestr.xlsx", _xlsx_bytes())
    assert [s.meta["sheet"] for s in segments] == ["Reklamacje"]  # empty sheet skipped
    assert segments[0].meta["header"] == "Nr | Data wpływu | Kwota | Status"
    assert "R-1 | 2026-03-02 | 1500 | uznana" in segments[0].text
    assert "R-2 | 2026-03-09 | 320.5 | odrzucona" in segments[0].text


def test_eml_headers_html_body_and_attachments():
    msg = EmailMessage()
    msg["From"] = "Jan Kowalski <jan.kowalski@example.com>"
    msg["To"] = "kadry@example.com"
    msg["Subject"] = "Wniosek urlopowy — załącznik"
    msg["Date"] = "Wed, 14 Oct 2026 09:30:00 +0200"
    msg.set_content("Tekst zwykły")
    msg.add_alternative("<p>Proszę o <b>akceptację</b> urlopu.</p><table><tr><td>A</td><td>B</td></tr></table>", subtype="html")
    msg.add_attachment(_docx_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename="umowa.docx")
    msg.add_attachment(b"\x00\x01", maintype="application", subtype="octet-stream", filename="dane.bin")
    text = as_text(extract_bytes("mail.eml", msg.as_bytes()))
    assert "From: Jan Kowalski <jan.kowalski@example.com>" in text and "Subject: Wniosek urlopowy" in text
    assert "Tekst zwykły" in text  # plain part preferred
    assert "Załączniki: umowa.docx, dane.bin" in text
    assert "--- Załącznik: umowa.docx ---" in text and "Przegląd | 1200 zł" in text


def test_html_to_text():
    assert html_to_text("<p>Ala</p><script>x=1</script><ul><li>a</li><li>b</li></ul>") == "Ala\n\na\n\nb"


def test_text_encodings():
    assert as_text(extract_bytes("a.txt", "Zażółć".encode("utf-8"))) == "Zażółć"
    assert as_text(extract_bytes("b.txt", "Zażółć".encode("cp1250"))) == "Zażółć"


def test_pdf_pages(tmp_path):
    from pathlib import Path

    from src.extractors import extract

    pdf = Path(__file__).resolve().parents[2] / "examples" / "demo_company" / "docs" / "cennik_uslug_2026.pdf"
    if not pdf.is_file():
        pytest.skip("demo PDF not generated yet")
    segments = extract(pdf)
    assert segments[0].meta["page"] == 1 and "Cennik" in segments[0].text


def test_unsupported_and_empty_files_are_errors():
    with pytest.raises(ExtractionError):
        extract_bytes("a.doc", b"x")
    with pytest.raises(ExtractionError):
        extract_bytes("empty.txt", b"   ")
    with pytest.raises(ExtractionError):
        extract_bytes("broken.docx", b"not a zip")
